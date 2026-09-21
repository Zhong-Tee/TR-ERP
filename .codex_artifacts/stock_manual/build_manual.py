from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"E:\Web_App\TR-ERP")
OUT = ROOT / "docs" / "warehouse-stock-verification-manual-th.docx"
IMAGES = {
    "warehouse_menu": Path(r"C:\Users\USER\AppData\Local\Temp\codex-clipboard-b829325a-855d-4957-be65-d5760fe56bac.png"),
    "sales": Path(r"C:\Users\USER\AppData\Local\Temp\codex-clipboard-b2dfbb03-b5b8-451d-9258-52afd8a3830d.png"),
    "subwarehouse": Path(r"C:\Users\USER\AppData\Local\Temp\codex-clipboard-6212d0cd-1f3a-43dc-ab54-97b038566f73.png"),
    "movement": Path(r"C:\Users\USER\AppData\Local\Temp\codex-clipboard-a1891b93-e9f1-4647-8704-870064b0af60.png"),
}

FONT = "Tahoma"
NAVY = "17365D"
BLUE = "2F65C8"
PALE_BLUE = "EAF2FC"
PALE_GRAY = "F5F7FA"
GRID = "D9D9D9"
TEXT = "202B3C"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=130, bottom=120, end=130):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text="", bold_lead=None, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, 10.5, True, TEXT)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, 10.5, False, TEXT)
    else:
        r = p.add_run(text)
        set_run_font(r, 10.5, False, TEXT)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    set_run_font(p.add_run(text), 10.3, False, TEXT)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    set_run_font(p.add_run(text), 10.5, False, TEXT)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        set_run_font(run, 15 if level == 1 else 12, True, "000000")
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr()
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Cm(widths[idx])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(header), 9.3, True, "FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx, value in enumerate(row):
            cell = cells[cidx]
            cell.width = Cm(widths[cidx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if ridx % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if cidx == 0 and len(headers) <= 4:
                set_run_font(p.add_run(str(value)), 9.2, True, TEXT)
            else:
                set_run_font(p.add_run(str(value)), 9.2, False, TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_image(doc, image_path, caption, width=6.75):
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    set_run_font(cp.add_run(caption), 8.8, False, "586577")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("หน้า ")
    set_run_font(r, 8.5, False, "5B6573")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r._r.append(fld_char1)
    r._r.append(instr)
    r._r.append(fld_char2)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.65)
section.bottom_margin = Cm(1.55)
section.left_margin = Cm(1.65)
section.right_margin = Cm(1.65)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10.5)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(7)
set_run_font(title.add_run("คู่มือตรวจสอบการตัดสต็อกและยอดคลังย่อย"), 22, True, "000000")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
set_run_font(subtitle.add_run("สำหรับตรวจสอบยอดขาย งาน WMS ความเคลื่อนไหวสต็อก และการใช้อะไหล่ในคลังย่อย"), 11.5, False, "46566A")

add_body(doc, "คู่มือนี้ใช้เมื่อต้องการตอบคำถามว่า สินค้าถูกตัดจากคลังหลักถูกต้องหรือไม่ และยอดใช้ในคลังย่อยคำนวณถูกต้องหรือไม่ หลักสำคัญคือแต่ละหน้าแสดงคนละเหตุการณ์และอาจใช้คนละวันที่ จึงต้องเทียบด้วยเลขบิล เลขใบงาน รหัสสินค้า และเวลาที่เกิดรายการ ไม่ใช่เทียบเฉพาะยอดรวมรายวัน")

add_heading(doc, "ภาพรวมของตัวเลขแต่ละหน้า", 1)
add_table(
    doc,
    ["หน้า", "ใช้ตรวจอะไร", "วันที่ที่ใช้", "ข้อควรจำ"],
    [
        ["รายการขายสินค้า", "สินค้าตามบิลที่จัดส่งแล้วหรือเสร็จสิ้น", "วันที่เปิดบิล", "เป็นข้อมูลเอกสารขาย ยังไม่ยืนยันว่าคลังตัดแล้ว"],
        ["จัดสินค้า ตรวจสินค้า", "รายการที่กำลังรอตรวจหรืออยู่ในกระบวนการ WMS", "เวลาทำรายการ WMS", "รายการที่ตรวจเสร็จแล้วอาจไม่อยู่ในหน้ารอตรวจ"],
        ["จัดสินค้า KPI", "ประวัติสรุปงาน WMS ที่ตรวจเสร็จ", "เวลาตรวจเสร็จ", "ใช้ตรวจย้อนหลังตามวันที่ปิดงาน"],
        ["รายการสินค้าคงเหลือ", "ยอดคงเหลือคลังหลักและความเคลื่อนไหวจริง", "เวลาที่เกิด movement", "เป็นหลักฐานสำคัญว่ามีการตัดหรือคืนสต็อกจริง"],
        ["คลังย่อย", "การใช้อะไหล่จากสินค้าผลิตที่จับคู่ไว้", "วันที่ปิดงาน WMS", "รหัสอะไหล่อาจไม่ใช่รหัสสินค้าที่ขาย"],
    ],
    [3.3, 5.2, 3.2, 5.3],
)

add_heading(doc, "เส้นทางเข้าเมนู WMS", 1)
add_body(doc, "หน้า คลัง และแท็บ คลังสินค้า ในแถบด้านบนเป็นหน้าสต็อกคลังหลัก ไม่ใช่หน้า WMS ดังนั้นจะไม่พบเมนู ตรวจสินค้า ในแถบนี้")
add_number(doc, "กดปุ่มเมนูสามขีดที่มุมซ้ายบน")
add_number(doc, "เลือกเมนูหลัก จัดสินค้า ซึ่งใช้เส้นทาง /wms")
add_number(doc, "เลือกเมนูย่อย ตรวจสินค้า สำหรับงานที่กำลังตรวจ หรือ KPI สำหรับค้นหางานย้อนหลัง")
add_image(doc, IMAGES["warehouse_menu"], "ภาพที่ 1 แถบเมนูคลังด้านบนไม่ใช่เมนู WMS ให้เปิดเมนูสามขีดและเลือก จัดสินค้า", 7.0)

doc.add_page_break()
add_heading(doc, "ขั้นตอนมาตรฐานในการตรวจหนึ่งกรณี", 1)

add_heading(doc, "ขั้นที่ 1 ระบุรายการที่จะตรวจ", 2)
add_body(doc, "จดข้อมูลต่อไปนี้ให้ครบก่อนเริ่ม เพื่อป้องกันการนำคนละบิลหรือคนละวันมาเทียบกัน")
for item in ["รหัสสินค้าสำเร็จรูป", "จำนวนในบิล", "เลขบิล", "ชื่อใบงาน", "วันที่เปิดบิล", "วันที่คลังตรวจ WMS"]:
    add_bullet(doc, item)

add_heading(doc, "ขั้นที่ 2 ตรวจรายการขายสินค้า", 2)
add_number(doc, "เข้า คลัง แล้วเลือก รายการขายสินค้า")
add_number(doc, "กำหนดวันที่ตามวันที่เปิดบิล และค้นหารหัสสินค้า")
add_number(doc, "กดชื่อสินค้าเพื่อดูรายการบิล ตรวจเลขบิล ชื่อใบงาน จำนวน และสถานะ")
add_body(doc, "ผลที่ได้จากขั้นนี้คือจำนวนตามเอกสารขาย ไม่ใช่หลักฐานว่ามี movement ตัดสต็อกแล้ว")
add_image(doc, IMAGES["sales"], "ภาพที่ 2 ตัวอย่างหน้ารายการขายสินค้า ซึ่งกรองด้วยวันที่เปิดบิล", 7.0)

add_heading(doc, "ขั้นที่ 3 ตรวจงาน WMS", 2)
add_number(doc, "เปิดเมนูสามขีด แล้วเข้า จัดสินค้า")
add_number(doc, "ถ้างานยังดำเนินการอยู่ ให้ดูเมนู ตรวจสินค้า")
add_number(doc, "ถ้างานเสร็จแล้ว ให้ดูเมนู KPI และเลือกช่วงวันที่ตรวจเสร็จ")
add_number(doc, "ยืนยันว่าจำนวนที่ตรวจถูกต้องตรงกับรายการที่ต้องหยิบ และไม่มีการยกเลิกหรือคืนคลังที่ยังไม่จบ")

doc.add_page_break()
add_heading(doc, "ขั้นที่ 4 ยืนยันการตัดสต็อกคลังหลัก", 2)
add_number(doc, "เข้า คลัง แล้วเลือก รายการสินค้าคงเหลือ")
add_number(doc, "ค้นหารหัสสินค้าสำเร็จรูป แล้วกดดู ความเคลื่อนไหว")
add_number(doc, "เลือกช่วงวันที่ตามวันที่ตรวจ WMS ไม่ใช่วันที่เปิดบิล")
add_number(doc, "กรองด้วยเลขใบงาน และดูเอกสารอ้างอิง งานคลัง WMS")
add_number(doc, "รวมรายการ ตัดสินค้า เบิกสินค้า แล้วหักรายการ คืนยอดจากการตัดสินค้า")
add_body(doc, "สูตรตรวจ WMS คือ จำนวนตัดสินค้า ลบ จำนวนคืนยอด เท่ากับ จำนวน WMS ที่ตรวจถูกต้อง", bold_lead="สูตรตรวจ WMS คือ ")
add_image(doc, IMAGES["movement"], "ภาพที่ 3 ตัวอย่างความเคลื่อนไหวสินค้า รายการติดลบคือการตัดสต็อกและมีใบงาน WMS อ้างอิง", 6.7)

add_body(doc, "ห้ามนำยอด จ่ายออก ทั้งหมดไปเทียบกับยอดขาย เพราะจ่ายออกอาจรวมใบเบิก ของเสีย การปรับสต็อก การเบิกเข้าผลิต และรายการที่ไม่เกี่ยวกับการขาย")

add_heading(doc, "ขั้นที่ 5 ตรวจคลังย่อย", 2)
add_number(doc, "เข้า คลัง แล้วเลือก คลังย่อย")
add_number(doc, "เลือกคลังย่อยและวันที่นับ")
add_number(doc, "ตรวจคอลัมน์ ใบงาน ใช้ตามใบงาน และคงเหลือสิ้นวัน")
add_number(doc, "หากมีสิทธิ์ตั้งค่า ให้เปิด ตั้งค่าหน้ายาง เพื่อตรวจว่าสินค้าผลิตถูกจับคู่กับอะไหล่รหัสใด")

add_body(doc, "คลังย่อยคำนวณจากสินค้าผลิตที่ตรวจ WMS แล้วกระจายยอดไปยังอะไหล่ทุกตัวในกลุ่มจับคู่ จึงไม่ควรเทียบรหัสอะไหล่กับรหัสสินค้าขายแบบหนึ่งต่อหนึ่งโดยไม่ตรวจ mapping")
add_image(doc, IMAGES["subwarehouse"], "ภาพที่ 4 ตัวอย่างคลังย่อย กลุ่ม WB แสดงการใช้อะไหล่ตามยอด WMS ของสินค้าผลิตที่จับคู่ไว้", 7.0)

doc.add_page_break()
add_heading(doc, "สูตรที่ใช้เทียบยอด", 1)
add_table(
    doc,
    ["จุดตรวจ", "สูตร", "ผลที่ควรได้"],
    [
        ["WMS กับคลังหลัก", "ตัดสินค้า - คืนยอด", "เท่ากับจำนวน WMS ที่ตรวจถูกต้อง"],
        ["ยอดคงเหลือคลังหลัก", "ยอดก่อนหน้า + movement ที่มีเครื่องหมายทั้งหมด", "เท่ากับยอดคงเหลือหลังรายการ"],
        ["คลังย่อยรายวัน", "ต้นวัน + เติม - ลดมือ - ใช้ตามใบงาน", "เท่ากับคงเหลือสิ้นวัน"],
        ["คลังย่อยช่วงวันที่", "เติม - ลดมือ - ใช้ตามใบงาน", "เท่ากับการเปลี่ยนแปลงสุทธิในช่วง"],
    ],
    [4.0, 7.0, 6.0],
)

add_heading(doc, "สาเหตุปกติที่ตัวเลขไม่ตรงกัน", 1)
add_table(
    doc,
    ["อาการ", "สาเหตุที่พบบ่อย", "วิธีตรวจ"],
    [
        ["ยอดขายกับ WMS คนละวัน", "บิลเปิดวันหยุด แต่คลังตรวจวันทำงานถัดไป", "ขยายช่วงวันที่ให้ครอบคลุมทั้งวันเปิดบิลและวันตรวจ แล้วเทียบด้วยเลขใบงาน"],
        ["ยอดจ่ายออกมากกว่ายอดขาย", "มีใบเบิก ของเสีย ปรับสต็อก หรือเบิกเข้าผลิต", "เลือกเฉพาะ movement ที่อ้างอิง งานคลัง WMS"],
        ["คลังย่อยมากกว่าสินค้าหนึ่งรหัส", "หนึ่งกลุ่มมีสินค้าผลิตหลายรหัส", "รวม WMS ของสินค้าผลิตทุกตัวในกลุ่ม mapping"],
        ["อะไหล่สองรหัสมียอดใช้เท่ากัน", "สินค้าหนึ่งชิ้นใช้ส่วนประกอบทั้งสองตัวในกลุ่ม", "ตรวจ ตั้งค่าหน้ายาง และรายชื่ออะไหล่ในกลุ่ม"],
        ["ยอดติดลบหรือไม่สมเหตุผล", "เติมสต็อกไม่ครบ ลดมือซ้ำ หรือ mapping ผิด", "ตรวจประวัติเติมลดและรายการ WMS ทีละใบงาน"],
        ["รายงานไม่มีใบงานแต่ movement มี", "ข้อมูลสรุป WMS เคยบันทึกไม่สำเร็จ", "ใช้ movement ยืนยันการตัดจริงและแจ้งผู้ดูแลระบบ ปัจจุบันมีเวลาปิดงานจากฐานข้อมูลสำรองแล้ว"],
    ],
    [4.0, 6.0, 7.0],
)

add_heading(doc, "ตัวอย่างกรณีกลุ่ม WB", 1)
add_body(doc, "กลุ่ม WB จับคู่สินค้าผลิต 110000085 110000086 และ 110000087 กับอะไหล่ 990000156 และ 990000222 เมื่อสินค้าผลิตในกลุ่มถูกตรวจ WMS หนึ่งชิ้น ระบบจะนับการใช้อะไหล่ทั้งสองรหัสตามกลุ่ม")
add_table(
    doc,
    ["รายการ", "จำนวน", "คำอธิบาย"],
    [
        ["รายการขายวันที่ 21", "7", "นับเฉพาะบิลที่มีวันที่เปิดบิลเป็นวันที่ 21"],
        ["คลังย่อยเดิม", "8", "รวมบิลที่เปิดวันที่ 20 แต่มาตรวจ WMS วันที่ 21 เพิ่มอีก 1"],
        ["ใบงาน SPTR 210969 R1", "3", "ตรวจวันที่ 21 แต่เดิมขาดจากรายงาน ต่อมาแก้ระบบและ backfill แล้ว"],
        ["ยอดใช้ WB หลังแก้", "11", "8 + 3 และถูกใช้กับอะไหล่ทั้ง 990000156 และ 990000222"],
    ],
    [5.0, 2.2, 9.8],
)

doc.add_page_break()
add_heading(doc, "แบบตรวจสอบก่อนสรุปผล", 1)
add_body(doc, "ใช้รายการนี้ทุกครั้งก่อนแจ้งว่าสต็อกผิด เพื่อแยกความต่างตามเวลา การยกเลิก และ mapping ออกจากความผิดพลาดจริง")
add_table(
    doc,
    ["ลำดับ", "รายการตรวจ", "ผลตรวจ"],
    [
        ["1", "เลขบิลและชื่อใบงานตรงกัน", "ผ่าน ไม่ผ่าน"],
        ["2", "จำนวนในบิลตรงกับจำนวนที่ควรหยิบ", "ผ่าน ไม่ผ่าน"],
        ["3", "สถานะ WMS เป็นถูกต้องหรือเป็นสถานะสุดท้ายที่เหมาะสม", "ผ่าน ไม่ผ่าน"],
        ["4", "มี movement ตัดสินค้าของใบงาน", "ผ่าน ไม่ผ่าน"],
        ["5", "หัก movement คืนยอดหรือยกเลิกแล้ว", "ผ่าน ไม่ผ่าน"],
        ["6", "ใช้ช่วงวันที่ครอบคลุมวันเปิดบิลและวันตรวจ WMS", "ผ่าน ไม่ผ่าน"],
        ["7", "ตรวจ mapping ระหว่างสินค้าผลิตและอะไหล่", "ผ่าน ไม่ผ่าน"],
        ["8", "สูตรต้นวัน เติม ลดมือ ใช้ตามใบงาน เท่ากับสิ้นวัน", "ผ่าน ไม่ผ่าน"],
    ],
    [1.5, 11.5, 4.0],
)

add_heading(doc, "ข้อมูลที่ควรส่งให้ผู้ดูแลระบบเมื่อยังไม่ตรง", 1)
for item in [
    "วันที่เปิดบิลและวันที่ตรวจ WMS",
    "เลขบิลและชื่อใบงาน",
    "รหัสสินค้าและจำนวนที่คาดหวัง",
    "ภาพหน้ารายการขายสินค้า",
    "ภาพความเคลื่อนไหวสินค้าที่เห็นเลขใบงาน",
    "ชื่อคลังย่อย รหัสอะไหล่ และยอดต้นวันกับสิ้นวัน",
    "ระบุว่ามีการยกเลิก คืนสินค้า ลดมือ หรือเติมสต็อกหรือไม่",
]:
    add_bullet(doc, item)

add_body(doc, "ข้อสรุป: หากต้องการยืนยันว่าคลังหลักตัดสต็อกจริง ให้ยึดความเคลื่อนไหวสินค้าที่อ้างอิงงานคลัง WMS เทียบกับจำนวน WMS เป็นหลัก รายการขายใช้ยืนยันเอกสาร ส่วนคลังย่อยใช้ยืนยันการใช้อะไหล่ตาม mapping", bold_lead="ข้อสรุป: ", after=0)

for sec in doc.sections:
    add_page_number(sec.footer.paragraphs[0])

doc.core_properties.title = "คู่มือตรวจสอบการตัดสต็อกและยอดคลังย่อย"
doc.core_properties.subject = "ขั้นตอนตรวจยอดขาย WMS ความเคลื่อนไหวสินค้า และคลังย่อย"
doc.core_properties.author = "TR ERP"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
